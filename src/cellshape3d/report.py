"""Stage 7: figures and tables from the stage-6 outputs (no statistic is recomputed here).

Figures (06_stats/figures/), journal style: 183 mm max width, Arial 5-7 pt, 0.5 pt axes, Okabe-Ito colours plus marker
shape as a redundant cue, vector PDF with embedded TrueType text and 600 dpi RGB PNG. One point = one image.
  Fig1_aspect_ratio     cell AR per group, one panel per cell line, Holm-adjusted ANOVA contrasts
  Fig2_alignment        a,b angle distributions (mean +/- s.e.m. of per-image histograms); c,d per-image median vs 45
  Fig3_relative_angle   signed object-minus-fibre angle histograms per fibre group, alignment index A = <cos 2 dtheta>
  Fig4_dtheta_vs_AR     dtheta distribution per AR class, per-image median |dtheta| vs AR
  source_data_*.csv     the plotted values
Tables (06_stats/): Table1_AR_summary.{csv,md[,docx]}, Table2_alignment_summary.{csv,md[,docx]}; median [IQR] across
images; superscript letters = compact letter display of the Holm-adjusted ANOVA contrasts.
"""
from __future__ import annotations

import re
from itertools import combinations
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from .config import Config  # noqa: E402
from .fibres import signed_rel_angle  # noqa: E402
from .stats import aligned_mask, compact_letters  # noqa: E402

MM = 1 / 25.4
OKABE_ITO = ["#0072B2", "#009E73", "#D55E00", "#CC79A7", "#E69F00", "#56B4E9", "#F0E442", "#000000"]
MARKERS = ["o", "s", "^", "D", "v", "P", "X", "*"]
STYLE = {
    "font.family": "sans-serif", "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"], "font.size": 6,
    "axes.labelsize": 6.5, "axes.titlesize": 6.5, "axes.linewidth": 0.5, "axes.spines.top": False,
    "axes.spines.right": False, "xtick.labelsize": 6, "ytick.labelsize": 6, "xtick.major.width": 0.5,
    "ytick.major.width": 0.5, "xtick.major.size": 2.5, "ytick.major.size": 2.5, "legend.fontsize": 6,
    "legend.frameon": False, "pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none",
    "savefig.facecolor": "white", "figure.facecolor": "white", "mathtext.default": "regular"}
JITTER = 0.13


class _Style:
    def __init__(self, cfg: Config):
        self.cfg = cfg
        gs = cfg.design.groups
        self.col = {g: OKABE_ITO[i % len(OKABE_ITO)] for i, g in enumerate(gs)}
        self.mrk = {g: MARKERS[i % len(MARKERS)] for i, g in enumerate(gs)}


def fmt_p(p: float) -> str:
    if not np.isfinite(p):
        return "n/a"
    if p < 1e-3:
        m, e = f"{p:.1e}".split("e")
        return rf"$\it{{P}}$ = {m} × 10$^{{{int(e)}}}$"
    return rf"$\it{{P}}$ = {p:.3f}" if p < 0.1 else rf"$\it{{P}}$ = {p:.2f}"


def _panel_label(ax, s):
    ax.text(-0.13, 1.04, s, transform=ax.transAxes, fontsize=8, fontweight="bold", va="bottom", ha="left")


def _strip_box(ax, x, v, g, st: _Style, rng):
    if len(v) == 0:
        return
    ax.boxplot(v, positions=[x], widths=0.56, showfliers=False, patch_artist=True, zorder=1,
               boxprops=dict(facecolor="none", edgecolor="black", lw=0.5), whiskerprops=dict(lw=0.5),
               capprops=dict(lw=0.5), medianprops=dict(color="black", lw=0.9))
    ax.scatter(x + rng.uniform(-JITTER, JITTER, len(v)), v, s=9, marker=st.mrk[g], facecolor=st.col[g],
               edgecolor="white", linewidth=0.3, alpha=0.9, zorder=3, clip_on=False)


def _bracket(ax, x1, x2, y, h, text):
    ax.plot([x1, x1, x2, x2], [y, y + h, y + h, y], lw=0.5, color="black", clip_on=False)
    ax.text((x1 + x2) / 2, y + h, text, ha="center", va="bottom", fontsize=5)


def _save(fig, D: Path, name: str):
    from PIL import Image
    fig.savefig(D / f"{name}.pdf")
    fig.savefig(D / f"{name}.png", dpi=600)
    plt.close(fig)
    # Agg writes RGBA; flatten to opaque RGB so journals see no alpha channel
    Image.open(D / f"{name}.png").convert("RGB").save(D / f"{name}.png", dpi=(600, 600))


# ------------------------------------------------------------------------------------------------ quick look
def quicklook_figures(nuc, cel, summ, cfg: Config, D: Path):
    """Quick-look plots written by the stats stage: AR by group and alignment histograms."""
    st = _Style(cfg)
    fig, ax = plt.subplots(1, 2, figsize=(10.5, 4.6))
    rng = np.random.default_rng(0)
    for A, (p, t) in zip(ax, [("nuc", "Nucleus AR"), ("cell", "Cell AR")]):
        s = summ[summ[f"{p}_ok"]]
        pos, ticks, labels = 0.0, [], []
        for line in cfg.design.cell_lines:
            for g in cfg.design.groups:
                v = s[(s.cell_line == line) & (s.group == g)][f"{p}_AR_median"].dropna().values
                if len(v):
                    A.boxplot(v, positions=[pos], widths=0.6, showfliers=False, medianprops=dict(color="k"))
                    A.scatter(np.full(len(v), pos) + rng.uniform(-0.15, 0.15, len(v)), v, s=18, color=st.col[g], zorder=3)
                ticks.append(pos)
                labels.append(f"{g}\n{line}")
                pos += 1
            pos += 0.6
        A.set_xticks(ticks)
        A.set_xticklabels(labels)
        A.set_title(t)
        A.set_ylabel("per-image median AR (xy)")
        A.spines[["top", "right"]].set_visible(False)
    fig.suptitle("Each point = one image; boxes summarise images")
    plt.tight_layout()
    for ext in ("png", "pdf"):
        plt.savefig(D / f"fig_AR_by_group.{ext}", dpi=200)
    plt.close(fig)
    fig, ax = plt.subplots(1, 2, figsize=(11, 4.2))
    for A, (df, t) in zip(ax, [(nuc[nuc.included], "Nuclei"), (cel[cel.included], "Cells")]):
        d = df[df.group.isin(cfg.design.fibre_groups) & aligned_mask(df, cfg)]
        for g in cfg.design.fibre_groups:
            v = d[d.group == g].align_deg
            if len(v):
                A.hist(v, bins=np.arange(0, 91, 10), density=True, histtype="step", lw=2, color=st.col[g],
                       label=f"{g} (n={len(v)})")
        A.axhline(1 / 90, color="gray", ls="--", lw=1, label="random")
        A.set_xlabel("angle between object long axis and local fibre (deg)")
        A.set_ylabel("density")
        A.set_title(f"{t} on fibres, AR > {cfg.stats.elong_ar}")
        A.legend(frameon=False)
        A.spines[["top", "right"]].set_visible(False)
    plt.tight_layout()
    for ext in ("png", "pdf"):
        plt.savefig(D / f"fig_alignment.{ext}", dpi=200)
    plt.close(fig)


# ------------------------------------------------------------------------------------------------ figures
def _pairs_staggered(groups):
    """All pairs, adjacent pairs first, so brackets can be stacked without overlapping."""
    return sorted(combinations(range(len(groups)), 2), key=lambda ij: (ij[1] - ij[0], ij[0]))


def fig1(summ, anova, cfg: Config, D: Path):
    st = _Style(cfg)
    groups = cfg.design.groups
    lines = cfg.report.panel_cell_lines or cfg.design.cell_lines
    metric = "cell AR (per-image median)"
    s = summ[summ["cell_ok"].astype(bool)].dropna(subset=["cell_AR_median"])
    s = s[s.cell_line.isin(lines)]
    if s.empty:
        return
    fig, axs = plt.subplots(1, len(lines), figsize=(60 * len(lines) * MM, 62 * MM), sharey=True,
                            layout="constrained", squeeze=False)
    axs = axs[0]
    rng = np.random.default_rng(cfg.report.seed)
    src = []
    allv = s.cell_AR_median.values
    span = max(allv.max() - allv.min(), 1e-3)
    h = 0.025 * span
    pairs = _pairs_staggered(groups)
    for ax, lab, l in zip(axs, "abcdefgh", lines):
        for x, g in enumerate(groups):
            d = s[(s.cell_line == l) & (s.group == g)]
            _strip_box(ax, x, d.cell_AR_median.values, g, st, rng)
            src += [dict(panel=f"1{lab}", metric=metric, cell_line=l, group=g, image=im, value=v)
                    for im, v in zip(d.image, d.cell_AR_median)]
        top = s[s.cell_line == l].cell_AR_median.max()
        t = anova[(anova.metric == metric) & (anova.cell_line == l)].set_index("term").p_holm if len(anova) else {}
        for k, (i, j) in enumerate(pairs):
            key = f"{groups[j]} vs {groups[i]}"
            if key in t:
                _bracket(ax, i, j, top + (0.06 + 0.12 * k) * span, h, fmt_p(t[key]))
        ax.set_xticks(range(len(groups)))
        ax.set_xticklabels(groups)
        ax.set_xlim(-0.6, len(groups) - 0.4)
        ax.set_title(l, fontweight="bold")
        _panel_label(ax, lab)
        ax.tick_params(labelleft=True)
    axs[0].set_ylabel("Cell aspect ratio")
    axs[0].set_ylim(allv.min() - 0.05 * span, allv.max() + (0.08 + 0.12 * len(pairs)) * span)
    _save(fig, D, "Fig1_aspect_ratio")
    pd.DataFrame(src).to_csv(D / "source_data_fig1.csv", index=False)


def _aligned(df, cfg):
    df = df[df.included.astype(bool) & df.group.isin(cfg.design.fibre_groups)]
    return df[aligned_mask(df, cfg)]


def fig2(nuc, cel, summ, tests, cfg: Config, D: Path):
    st = _Style(cfg)
    rep, nmin = cfg.report, cfg.stats.min_objects_per_image
    fg, lines = cfg.design.fibre_groups, cfg.design.cell_lines
    fig, axs = plt.subplots(2, 2, figsize=(183 * MM, 112 * MM), layout="constrained")
    edges = np.arange(0, 90 + rep.bin_deg, rep.bin_deg)
    ctr = edges[:-1] + rep.bin_deg / 2
    offs = np.linspace(-1.2, 1.2, len(fg)) if len(fg) > 1 else [0.0]
    hsrc, msrc = [], []
    for ax, lab, (df, name, adj) in zip(axs[0], "ab", [(nuc, "Nuclei", "Nuclear"), (cel, "Cells", "Cell")]):
        d = _aligned(df, cfg)
        for g, off, ls in zip(fg, offs, ["-", "--", ":", "-."] * 3):
            h = []
            for im, dd in d[d.group == g].groupby("image"):
                if len(dd) < nmin:
                    continue
                f = np.histogram(dd.align_deg.clip(0, 90), edges)[0] / len(dd) * 100
                h.append(f)
                hsrc += [dict(panel=f"2{lab}", objects=name, group=g, image=im, bin_lo=e0, bin_hi=e0 + rep.bin_deg,
                              percent=x, n_objects=len(dd)) for e0, x in zip(edges[:-1], f)]
            if not h:
                continue
            h = np.array(h)
            se = h.std(0, ddof=1) / np.sqrt(len(h)) if len(h) > 1 else np.zeros(len(ctr))
            n_obj = int(d[d.group == g].groupby("image").filter(lambda x: len(x) >= nmin).shape[0])
            ax.errorbar(ctr + off, h.mean(0), se, color=st.col[g], marker=st.mrk[g], ms=3, lw=0.9, ls=ls,
                        elinewidth=0.6, capsize=1.2, mec="white", mew=0.3,
                        label=f"{g} ({len(h)} images, {n_obj:,} {name.lower()})")
        ax.axhline(100 / len(ctr), color="0.5", lw=0.6, ls=":", zorder=0)
        ax.text(89, 100 / len(ctr) + 0.6, "Random", ha="right", va="bottom", fontsize=5.5, color="0.4")
        ax.set_xlim(0, 90)
        ax.set_xticks(range(0, 91, 15))
        ax.set_ylim(0, None)
        ax.set_xlabel(f"{adj} long axis–fibre angle (°)")
        ax.set_ylabel(f"{name} (% per image)")
        if ax.get_legend_handles_labels()[0]:
            ax.legend(loc="upper right", handlelength=2.2)
        _panel_label(ax, lab)
    pos = {(l, g): i + j * (len(fg) + 0.6) for j, l in enumerate(lines) for i, g in enumerate(fg)}
    rng = np.random.default_rng(rep.seed)
    for ax, lab, (p, name) in zip(axs[1], "cd", [("nuc", "Nuclear"), ("cell", "Cell")]):
        col = f"{p}_align_elong_median_deg"
        t = tests[tests.metric == f"{p} alignment to fibre (elongated, median deg)"] if len(tests) else pd.DataFrame()
        for (l, g), x in pos.items():
            s = summ[(summ.cell_line == l) & (summ.group == g)].dropna(subset=[col]) if col in summ else summ.iloc[:0]
            v = s[col].values if col in s else np.array([])
            _strip_box(ax, x, v, g, st, rng)
            if len(t):
                pv = t[(t.cell_line == l) & (t.test == f"Wilcoxon {g}: median vs 45 deg")].p
                if len(pv):
                    ax.text(x, 54, fmt_p(pv.iloc[0]), ha="center", va="bottom", fontsize=5)
            msrc += [dict(panel=f"2{lab}", objects=p, cell_line=l, group=g, image=im, median_deg=val, n_objects=n)
                     for im, val, n in zip(s.image, v, s.get(f"{p}_align_elong_n", []))]
        ax.axhline(45, color="0.5", lw=0.6, ls=":", zorder=0)
        ax.text(max(pos.values()) + 0.55, 45.8, "Random (45°)", ha="right", va="bottom", fontsize=5.5, color="0.4")
        ax.set_ylim(0, 62)
        ax.set_yticks(range(0, 61, 15))
        ax.set_xticks(list(pos.values()))
        ax.set_xticklabels([g for (_, g) in pos])
        ax.set_xlim(min(pos.values()) - 0.6, max(pos.values()) + 0.6)
        for l in lines:
            xs = [pos[(l, g)] for g in fg]
            ax.annotate(l, xy=(np.mean(xs), 0), xycoords=("data", "axes fraction"), xytext=(0, -15),
                        textcoords="offset points", ha="center", va="top", fontsize=6.5)
        ax.set_ylabel(f"{name} long axis–fibre angle\n(per-image median, °)")
        _panel_label(ax, lab)
    _save(fig, D, "Fig2_alignment")
    pd.DataFrame(hsrc).to_csv(D / "source_data_fig2_hist.csv", index=False)
    pd.DataFrame(msrc).to_csv(D / "source_data_fig2_median.csv", index=False)


def fig3(cel, cfg: Config, D: Path, log=print):
    from scipy.stats import wilcoxon
    rep, nmin, fg = cfg.report, cfg.stats.min_objects_per_image, cfg.design.fibre_groups
    d = _aligned(cel, cfg).copy()
    d["rel_deg"] = signed_rel_angle(d.orientation_deg, d.fibre_angle_deg)
    edges = np.arange(-90, 90 + rep.bin_deg, rep.bin_deg)
    ctr = edges[:-1] + rep.bin_deg / 2
    fig, axs = plt.subplots(1, len(fg), figsize=(60 * len(fg) * MM, 58 * MM), sharey=True, layout="constrained",
                            squeeze=False)
    axs = axs[0]
    hsrc, ssrc, ytop = [], [], 1.0
    for ax, lab, g in zip(axs, "abcdefgh", fg):
        h, S = [], []
        for im, dd in d[d.group == g].groupby("image"):
            if len(dd) < nmin:
                continue
            f = np.histogram(dd.rel_deg, edges)[0] / len(dd) * 100
            h.append(f)
            s_im = np.cos(np.radians(2 * dd.rel_deg)).mean()
            S.append(s_im)
            hsrc += [dict(panel=f"3{lab}", group=g, cell_line=dd.cell_line.iloc[0], image=im, bin_lo=e0,
                          bin_hi=e0 + rep.bin_deg, percent=x, n_cells=len(dd)) for e0, x in zip(edges[:-1], f)]
            ssrc.append(dict(group=g, cell_line=dd.cell_line.iloc[0], image=im, alignment_A=s_im, n_cells=len(dd)))
        ax.set_title(g, fontweight="bold")
        _panel_label(ax, lab)
        ax.set_xlim(-90, 90)
        ax.set_xticks(range(-90, 91, 30))
        ax.set_xlabel("Cell–fibre angle (°)")
        if not h:
            continue
        h, S = np.array(h), np.array(S)
        m = h.mean(0)
        se = h.std(0, ddof=1) / np.sqrt(len(h)) if len(h) > 1 else np.zeros(len(ctr))
        ytop = max(ytop, (m + se).max())
        pw = wilcoxon(S).pvalue if len(S) >= 3 else np.nan
        log(f"Fig3 {g}: {len(h)} images, A = {S.mean():.2f} (n={len(S)}), Wilcoxon vs 0 P = {pw:.2g}")
        ax.bar(ctr, m, width=rep.bin_deg * 0.86, color="0.62", edgecolor="none", zorder=2)
        ax.errorbar(ctr, m, se, fmt="none", ecolor="black", elinewidth=0.5, capsize=1.2, capthick=0.5, zorder=3)
        ax.axhline(100 / len(ctr), color="black", lw=0.6, ls=(0, (3, 2)), zorder=4)
    axs[0].set_ylim(0, ytop * 1.08)
    axs[0].set_ylabel("Cells (%)")
    for ax in axs:
        ax.tick_params(labelleft=True)
    _save(fig, D, "Fig3_relative_angle")
    pd.DataFrame(hsrc).to_csv(D / "source_data_fig3_hist.csv", index=False)
    pd.DataFrame(ssrc).to_csv(D / "source_data_fig3_alignment.csv", index=False)


AR_EDGES = [1.0, 1.25, 1.5, 2.0, 2.5, 3.0, np.inf]
AR_LABELS = ["1–1.25", "1.25–1.5", "1.5–2", "2–2.5", "2.5–3", "≥3"]


def fig4(cel, cfg: Config, D: Path, log=print):
    """dtheta vs cell AR (same cells as Fig. 3 but without the AR cut-off, since AR is the x variable)."""
    from scipy.stats import spearmanr, wilcoxon
    st = _Style(cfg)
    rep, nmin, fg = cfg.report, cfg.stats.min_objects_per_image, cfg.design.fibre_groups
    d = cel[cel.included.astype(bool) & cel.group.isin(fg)]
    d = d[aligned_mask(d, cfg, elongated=False)].copy()
    if d.empty:
        return
    d["rel_deg"] = signed_rel_angle(d.orientation_deg, d.fibre_angle_deg)
    d["ar_bin"] = pd.cut(d.AR_xy, AR_EDGES, right=False, labels=False)
    th_edges = np.arange(-90, 90 + rep.bin_deg, rep.bin_deg)
    nk = len(AR_LABELS)
    top = ";".join(["".join("abcdefgh"[:len(fg)]), "".join(["z"] * len(fg))])
    fig, axs = plt.subplot_mosaic(top, figsize=(60 * len(fg) * MM, 105 * MM), layout="constrained",
                                  gridspec_kw=dict(height_ratios=[1.15, 1]))
    dens, hsrc, ssrc, rsrc = {}, [], [], []
    for lab, g in zip("abcdefgh", fg):
        dg = d[d.group == g]
        H = np.array([np.histogram(dg[dg.ar_bin == k].rel_deg, th_edges)[0] for k in range(nk)]).T.astype(float)
        dens[g] = H / np.maximum(H.sum(0), 1) * 100       # % of cells within each AR class
        hsrc += [dict(panel=f"4{lab}", group=g, ar_class=AR_LABELS[k], dtheta_lo=th_edges[i], dtheta_hi=th_edges[i + 1],
                      n_cells=int(H[i, k]), percent_of_ar_class=dens[g][i, k])
                 for i in range(len(th_edges) - 1) for k in range(nk)]
    vmax = max(max(v.max() for v in dens.values()), 1e-6)
    im = None
    for lab, g in zip("abcdefgh", fg):
        ax = axs[lab]
        im = ax.pcolormesh(np.arange(nk + 1) - 0.5, th_edges, dens[g], cmap="cividis", vmin=0, vmax=vmax,
                           shading="flat", rasterized=True)
        ax.set_xticks(np.arange(nk) - 0.5)
        ax.set_xticklabels(["1", "1.25", "1.5", "2", "2.5", "3"])
        ax.set_yticks(range(-90, 91, 45))
        ax.set_ylabel("Δθ (°)" if lab == "a" else "")
        ax.set_xlabel("Cell AR")
        ax.set_title(g, fontweight="bold")
        _panel_label(ax, lab)
        for sp in ax.spines.values():
            sp.set_visible(False)
        ax.tick_params(length=0)
    cb = fig.colorbar(im, ax=[axs[k] for k in "abcdefgh"[:len(fg)]], shrink=0.85, aspect=18, pad=0.02)
    cb.set_label("Cells (%)")
    cb.outline.set_linewidth(0.5)
    ax = axs["z"]
    rng = np.random.default_rng(rep.seed)
    offs = np.linspace(-0.18, 0.18, len(fg)) if len(fg) > 1 else [0.0]
    for g, off in zip(fg, offs):
        dg = d[d.group == g]
        per = []
        for im_, di in dg.groupby("image"):
            if len(di) >= nmin and di.AR_xy.nunique() > 1:
                rsrc.append(dict(group=g, cell_line=di.cell_line.iloc[0], image=im_, n_cells=len(di),
                                 spearman_rho_AR_vs_abs_dtheta=spearmanr(di.AR_xy, di.rel_deg.abs()).statistic))
            for k, dk in di.groupby("ar_bin"):
                if len(dk) >= 5:
                    per.append(dict(group=g, cell_line=dk.cell_line.iloc[0], image=im_, ar_class=AR_LABELS[int(k)],
                                    k=int(k), median_abs_dtheta=dk.rel_deg.abs().median(),
                                    alignment_A=np.cos(np.radians(2 * dk.rel_deg)).mean(), n_cells=len(dk)))
        if not per:
            continue
        per = pd.DataFrame(per)
        ssrc.append(per)
        agg = per.groupby("k").median_abs_dtheta.agg(["mean", "sem", "count"])
        ax.scatter(per.k + off + rng.uniform(-0.07, 0.07, len(per)), per.median_abs_dtheta, s=7, marker=st.mrk[g],
                   facecolor=st.col[g], edgecolor="white", linewidth=0.3, alpha=0.85, zorder=2, label=g)
        ax.errorbar(agg.index + off, agg["mean"], agg["sem"].fillna(0), fmt="none", ecolor="black", elinewidth=0.6,
                    capsize=1.5, capthick=0.6, zorder=3)
        ax.hlines(agg["mean"], agg.index + off - 0.11, agg.index + off + 0.11, color="black", lw=0.9, zorder=3)
        r = pd.Series([x["spearman_rho_AR_vs_abs_dtheta"] for x in rsrc if x["group"] == g], dtype=float).dropna()
        if len(r) >= 3:
            log(f"Fig4 {g}: per-image Spearman rho (AR vs |dtheta|) = {r.mean():.2f} +/- {r.sem():.2f} "
                f"(n={len(r)}), Wilcoxon vs 0 P = {wilcoxon(r).pvalue:.2g}")
    ax.axhline(45, color="0.35", lw=0.6, ls=(0, (1.5, 1.5)), zorder=0)
    ax.set_xticks(range(nk))
    ax.set_xticklabels(AR_LABELS)
    ax.set_xlim(-0.6, nk - 0.4)
    ax.set_ylim(0, 90)
    ax.set_yticks(range(0, 91, 15))
    ax.set_xlabel("Cell AR")
    ax.set_ylabel("|Δθ| (°)")
    if ax.get_legend_handles_labels()[0]:
        ax.legend(loc="upper right", handletextpad=0.2, markerscale=1.3)
    _panel_label(ax, "abcdefgh"[len(fg)])
    _save(fig, D, "Fig4_dtheta_vs_AR")
    pd.DataFrame(hsrc).to_csv(D / "source_data_fig4_density.csv", index=False)
    if ssrc:
        pd.concat(ssrc).drop(columns="k").to_csv(D / "source_data_fig4_alignment.csv", index=False)
    pd.DataFrame(rsrc).to_csv(D / "source_data_fig4_spearman.csv", index=False)


def _read_stats(cfg: Config):
    S6 = Path(cfg.data.out_dir) / "06_stats"
    if not (S6 / "per_image_summary.csv").exists():
        raise FileNotFoundError(f"{S6 / 'per_image_summary.csv'} not found; run the stats stage first")

    def rd(name):
        try:
            return pd.read_csv(S6 / name)
        except pd.errors.EmptyDataError:
            return pd.DataFrame()

    use = ["image", "group", "cell_line", "orientation_deg", "fibre_angle_deg", "included", "on_fibre",
           "fibre_coherence", "AR_xy", "align_deg", "proj_area_um2"]
    nuc = pd.read_csv(S6 / "all_nuclei.csv", usecols=lambda c: c in use, low_memory=False)
    cel = pd.read_csv(S6 / "all_cells.csv", usecols=lambda c: c in use, low_memory=False)
    return S6, rd("per_image_summary.csv"), rd("group_tests.csv"), rd("anova_tests.csv"), nuc, cel


def make_figures(cfg: Config, log=print):
    S6, summ, tests, anova, nuc, cel = _read_stats(cfg)
    D = cfg.stage_dir("06_stats/figures")
    with plt.rc_context(STYLE):
        fig1(summ, anova, cfg, D)
        if "align_deg" in cel and cfg.design.fibre_groups:
            fig2(nuc, cel, summ, tests, cfg, D)
            fig3(cel, cfg, D, log)
            fig4(cel, cfg, D, log)
    log(f"figures written to {D}")


# ------------------------------------------------------------------------------------------------ tables
SECTION = "§"     # rows whose label starts with this are section headings spanning the table


def miqr(v, fmt: str) -> str:
    v = pd.Series(v, dtype=float).dropna()
    if v.empty:
        return "–"
    q1, q2, q3 = np.percentile(v, [25, 50, 75])
    return f"{fmt.format(q2)} [{fmt.format(q1)}–{fmt.format(q3)}]"


def fmt_p_text(p: float) -> str:
    if not np.isfinite(p):
        return "–"
    if p < 1e-3:
        m, e = f"{p:.1e}".split("e")
        return f"{m} × 10^{int(e)}^"
    return f"{p:.3f}"


def _table_frame(rows, cols):
    return pd.DataFrame([{"Variable": lab, **{f"{l} {g}": v for (l, g), v in zip(cols, vals)}} for lab, vals in rows])


def table1(summ, anova, nuc, cel, cfg: Config):
    groups, lines, alpha = cfg.design.groups, cfg.design.cell_lines, cfg.stats.alpha
    cols = [(l, g) for l in lines for g in groups]
    extra = pd.DataFrame({"nuc_area": nuc[nuc.included.astype(bool)].groupby("image").proj_area_um2.median(),
                          "cell_area": cel[cel.included.astype(bool)].groupby("image").proj_area_um2.median()})
    summ = summ.join(extra, on="image")
    cld = {}
    for p, metric in [("nuc", "nucleus AR (per-image median)"), ("cell", "cell AR (per-image median)")]:
        for line in lines:
            c = anova[(anova.metric == metric) & (anova.cell_line == line)] if len(anova) else pd.DataFrame()
            sig = {frozenset(t.split(" vs ")) for t, ph in zip(c.get("term", []), c.get("p_holm", [])) if ph < alpha}
            s = summ[summ[f"{p}_ok"].astype(bool) & (summ.cell_line == line)]
            med = s.groupby("group")[f"{p}_AR_median"].median()
            order = sorted(groups, key=lambda g: med.get(g, np.inf))
            for g, L in compact_letters(groups, sig, order).items():
                cld[(p, line, g)] = L if len(c) else ""
    sel = lambda l, g, ok: summ[(summ.cell_line == l) & (summ.group == g) & summ[ok].astype(bool)]
    R = lambda lab, fn: (lab, [fn(l, g) for l, g in cols])
    sup = lambda s: f"^{s}^" if s else ""
    rows = [R("Images (constructs), n", lambda l, g: f"{len(summ[(summ.cell_line == l) & (summ.group == g)])}"),
            R("Nuclei analysed, n", lambda l, g: f"{int(sel(l, g, 'nuc_ok').nuc_n.sum()):,}"),
            R("Cells analysed, n", lambda l, g: f"{int(sel(l, g, 'cell_ok').cell_n.sum()):,}")]
    for p, lab in [("nuc", "Nuclear AR"), ("cell", "Cell AR")]:
        rows.append(R(lab, lambda l, g, p=p: miqr(sel(l, g, f"{p}_ok")[f"{p}_AR_median"], "{:.2f}") + sup(cld[(p, l, g)])))
    rows += [R("Cells with AR > 2 (%)", lambda l, g: miqr(sel(l, g, "cell_ok").cell_frac_AR_gt2 * 100, "{:.1f}")),
             R("Nuclear projected area (µm²)", lambda l, g: miqr(sel(l, g, "nuc_ok").nuc_area, "{:.0f}")),
             R("Cell projected area (µm²)", lambda l, g: miqr(sel(l, g, "cell_ok").cell_area, "{:.0f}"))]
    note = ("Values are median [interquartile range] across images; each image is one replicate, summarised by the "
            "median of its objects. AR, aspect ratio (major/minor axis of the moment-equivalent ellipse of the xy "
            f"projection). AR rows include images with ≥{cfg.stats.min_objects_per_image} analysed objects. "
            "Superscript letters: within each cell line, groups that share a letter do not differ (Holm-adjusted "
            f"contrasts of a two-way ANOVA on log AR, α = {alpha}).")
    if len(anova):
        a = anova[anova.test.str.startswith("two-way")].set_index(["metric", "term"])
        parts = []
        for k, m in [("nuclear", "nucleus AR (per-image median)"), ("cell", "cell AR (per-image median)")]:
            if (m, "scaffold") in a.index:
                r = a.loc[(m, "scaffold")]
                parts.append(f"{k} AR *F*({int(r.df)},{int(r.df_resid)}) = {r.F:.1f}, *P* = {fmt_p_text(r.p)}")
        if parts:
            note += " Group effect: " + "; ".join(parts) + "."
    title = "Table 1 | Aspect ratio and size of cells and nuclei per group and cell line."
    return _table_frame(rows, cols), cols, title, note


def table2(summ, tests, nuc, cel, cfg: Config):
    s_cfg = cfg.stats
    groups, lines = cfg.design.fibre_groups, cfg.design.cell_lines
    cols = [(l, g) for l in lines for g in groups]
    extra = {}
    for p, df in [("nuc", nuc), ("cell", cel)]:
        df = df[df.included.astype(bool)]
        al = df[aligned_mask(df, cfg)]
        extra[f"{p}_pct_aligned"] = al.groupby("image").align_deg.apply(lambda v: (v <= s_cfg.aligned_deg).mean() * 100)
    summ = summ.join(pd.DataFrame(extra), on="image")
    sel = lambda l, g: summ[(summ.cell_line == l) & (summ.group == g)]
    R = lambda lab, fn: (lab, [fn(l, g) for l, g in cols])
    rows = [R("Images (constructs), n", lambda l, g: f"{len(sel(l, g))}")]
    for p, obj in [("nuc", "Nuclei"), ("cell", "Cells")]:
        wt = tests[tests.metric == f"{p} alignment to fibre (elongated, median deg)"] if len(tests) else pd.DataFrame()

        def pval(l, g, wt=wt):
            if wt.empty:
                return "–"
            v = wt[(wt.cell_line == l) & (wt.test == f"Wilcoxon {g}: median vs 45 deg")].p
            return fmt_p_text(v.iloc[0]) if len(v) else "–"

        rows += [(SECTION + obj, [""] * len(cols)),
                 R(f"Elongated {obj.lower()} on fibres, n",
                   lambda l, g, p=p: f"{int(sel(l, g)[f'{p}_align_elong_n'].fillna(0).sum()):,}"),
                 R("Long axis–fibre angle (°)", lambda l, g, p=p: miqr(sel(l, g)[f"{p}_align_elong_median_deg"], "{:.1f}")),
                 R(f"Within {s_cfg.aligned_deg:g}° of fibre axis (%)",
                   lambda l, g, p=p: miqr(sel(l, g)[f"{p}_pct_aligned"], "{:.1f}")),
                 R("Angle vs 45°, *P*", pval)]
    note = ("Values are median [interquartile range] across images. Only elongated objects "
            f"(AR > {s_cfg.elong_ar}) within {cfg.fibres.on_fibre_um:g} µm of a fibre, at positions with "
            f"structure-tensor coherence ≥ {s_cfg.align_coherence_min}, were analysed. Angle: per-image median of the "
            "angle between the object long axis and the local fibre axis (0° = parallel, 45° = random). "
            f"Within {s_cfg.aligned_deg:g}°: per-image percentage of objects with an angle ≤ {s_cfg.aligned_deg:g}° "
            f"({s_cfg.aligned_deg / 90 * 100:.1f}% expected for random orientation). *P*: two-sided one-sample "
            "Wilcoxon signed-rank test of per-image median angles against 45° (uncorrected).")
    title = "Table 2 | Alignment of nuclei and cells with the local fibre direction."
    return _table_frame(rows, cols), cols, title, note


def write_md(t, cols, title, note, path: Path):
    head = [l if i == 0 or cols[i - 1][0] != l else "" for i, (l, _) in enumerate(cols)]
    lines = [f"**{title}**", "", "| | " + " | ".join(head) + " |", "|" + "---|" * (len(cols) + 1),
             "| **Variable** | " + " | ".join(f"**{g}**" for _, g in cols) + " |"]
    for _, r in t.iterrows():
        v = r.Variable
        lines.append(f"| **{v[1:]}** |" + " |" * len(cols) if v.startswith(SECTION)
                     else "| " + " | ".join(str(r[c]) for c in t.columns) + " |")
    path.write_text("\n".join(lines + ["", note]), encoding="utf-8")


def _rich(par, text, size, bold=False):
    """Add text to a python-docx paragraph, turning ^x^ into superscript and *x* into italics."""
    from docx.shared import Pt
    for tok in re.split(r"(\^[^^]+\^|\*[^*]+\*)", text):
        if not tok:
            continue
        r = par.add_run(tok.strip("^*") if tok[0] in "^*" else tok)
        r.font.size = Pt(size)
        r.bold = bold
        r.font.superscript = tok.startswith("^")
        r.italic = tok.startswith("*")


def write_docx(t, cols, title, note, path: Path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(8)
    _rich(doc.add_paragraph(), title, 9, bold=True)
    tab = doc.add_table(rows=2 + len(t), cols=len(cols) + 1)
    tab.autofit = False
    w0 = 4.8
    widths = [Cm(w0)] + [Cm((17.4 - w0) / len(cols))] * len(cols)

    def border(cell, sides):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for s in sides:
            e = OxmlElement(f"w:{s}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "6")
            b.append(e)
        tcPr.append(b)

    h0, h1 = tab.rows[0].cells, tab.rows[1].cells
    for l in dict.fromkeys(c[0] for c in cols):
        idx = [i + 1 for i, c in enumerate(cols) if c[0] == l]
        c = h0[idx[0]].merge(h0[idx[-1]])
        c.text = ""
        _rich(c.paragraphs[0], l, 8, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, lab in enumerate(["Variable"] + [g for _, g in cols]):
        _rich(h1[k].paragraphs[0], lab, 8, bold=True)
        if k:
            h1[k].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (_, row) in enumerate(t.iterrows()):
        cells = tab.rows[i + 2].cells
        if row.Variable.startswith(SECTION):
            c = cells[0].merge(cells[-1])
            c.text = ""
            _rich(c.paragraphs[0], row.Variable[1:], 7.5, bold=True)
            continue
        for k, col in enumerate(t.columns):
            _rich(cells[k].paragraphs[0], str(row[col]), 7.5)
            if k:
                cells[k].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in tab.rows:
        for k, c in enumerate(row.cells):
            c.width = widths[k]
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.space_before = Pt(1)
    for c in h0:
        border(c, ["top"])
    for c in h1:
        border(c, ["bottom"])
    for c in tab.rows[-1].cells:
        border(c, ["bottom"])
    _rich(doc.add_paragraph(), note, 7)
    doc.save(path)


def make_tables(cfg: Config, log=print):
    S6, summ, tests, anova, nuc, cel = _read_stats(cfg)
    try:
        import docx  # noqa: F401
        have_docx = True
    except ImportError:
        have_docx = False
        log("python-docx not installed: writing .csv/.md only (pip install cellshape3d[docx])")
    out = [("Table1_AR_summary", table1(summ, anova, nuc, cel, cfg))]
    if "align_deg" in cel and cfg.design.fibre_groups:
        out.append(("Table2_alignment_summary", table2(summ, tests, nuc, cel, cfg)))
    for name, (t, cols, title, note) in out:
        c = t.copy()
        c["Variable"] = c.Variable.str.lstrip(SECTION)
        c.to_csv(S6 / f"{name}.csv", index=False, encoding="utf-8-sig")
        write_md(t, cols, title, note, S6 / f"{name}.md")
        if have_docx:
            write_docx(t, cols, title, note, S6 / f"{name}.docx")
        log(f"written {name}")
